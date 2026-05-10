from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

class ContractGenerator:
    def __init__(self, data):
        self.data = data
        self.doc = Document()

    def _add_heading(self, text):
        heading = self.doc.add_heading(text, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_paragraph(self, text, bold_parts=None, italic=False):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(12)
        if italic:
            run.italic = True
        return p

    def generate_lease_contract(self, filename="contract.docx"):
        # Налаштування шрифту за замовчуванням
        style = self.doc.styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(12)

        # 1. Заголовок
        self._add_heading("ДОГОВІР ОРЕНДИ ЖИТЛОВОГО ПРИМІЩЕННЯ")
        
        # 2. Дата та місце
        p_info = self.doc.add_paragraph()
        p_info.add_run(f"м. {self.data.get('city', '__________')}")
        p_info.add_run("\t" * 5)
        p_info.add_run(f"«___» ________ 202_ р.")

        # 3. Сторони
        self.doc.add_paragraph().add_run("1. СТОРОНИ ДОГОВОРУ").bold = True
        self._add_paragraph(
            f"Орендодавець: {self.data.get('landlord_name', '____________________')}, з однієї сторони, та\n"
            f"Орендар: {self.data.get('tenant_name', '____________________')}, з іншої сторони, "
            "уклали цей Договір про наступне:"
        )

        # 4. Предмет договору
        self.doc.add_paragraph().add_run("2. ПРЕДМЕТ ДОГОВОРУ").bold = True
        self._add_paragraph(
            f"2.1. Орендодавець передає, а Орендар приймає у строкове платне володіння та користування житлове приміщення "
            f"за адресою: {self.data.get('address', '____________________')}, загальною площею {self.data.get('area', '___')} кв.м."
        )

        # 5. Оплата
        self.doc.add_paragraph().add_run("3. ОРЕНДНА ПЛАТА ТА РОЗРАХУНКИ").bold = True
        self._add_paragraph(
            f"3.1. Розмір орендної плати за місяць становить {self.data.get('price', '_______')} грн.\n"
            "3.2. Орендар зобов'язується сплачувати оренду щомісяця не пізніше ___ числа поточного місяця."
        )

        # 6. Права та обов'язки (Статичний блок)
        self.doc.add_paragraph().add_run("4. ПРАВА ТА ОБОВ'ЯЗКИ СТОРІН").bold = True
        self._add_paragraph(
            "4.1. Орендар зобов’язується використовувати приміщення лише за призначенням.\n"
            "4.2. Орендар не має права передавати приміщення в суборенду без згоди Орендодавця."
        )

        # 7. Підписи сторін (Таблиця для зручності)
        self.doc.add_paragraph().add_run("\n5. РЕКВІЗИТИ ТА ПІДПИСИ СТОРІН").bold = True
        table = self.doc.add_table(rows=1, cols=2)
        cells = table.rows[0].cells
        
        cells[0].text = f"ОРЕНДОДАВЕЦЬ:\n\n________________\n(підпис)\n{self.data.get('landlord_name', '')}"
        cells[1].text = f"ОРЕНДАР:\n\n________________\n(підпис)\n{self.data.get('tenant_name', '')}"

        # Збереження
        self.doc.save(filename)
        return filename

# --- ПРИКЛАД ВИКОРИСТАННЯ ---
